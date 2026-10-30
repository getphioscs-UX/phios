import sys,json,io,zipfile,re
import pdfplumber
from openpyxl import load_workbook
b=sys.stdin.buffer.read();kind=sys.argv[1]
try:
 if kind=='pdf':
  if not b.startswith(b'%PDF-'):raise ValueError('PDF_SIGNATURE_INVALID')
  if b'/JavaScript' in b or b'/EmbeddedFile' in b or b'/OpenAction' in b:raise ValueError('SUSPICIOUS_PDF_QUARANTINED')
  out=[]
  with pdfplumber.open(io.BytesIO(b)) as pdf:
   if len(pdf.pages)>30:raise ValueError('PDF_PAGE_LIMIT')
   for i,p in enumerate(pdf.pages):
    text=p.extract_text() or ''
    for line in text.splitlines():
     m=re.fullmatch(r'(\d{4}-Q[1-4]) \| (\w+) \| (-?\d+\.\d{2}) \| ([A-Z]{3})',line)
     if m:out.append({'period':m[1],'field':m[2],'rawValue':m[3],'currency':m[4],'sourceLocation':{'page':i+1,'text':line}})
  if not out:raise ValueError('TEXT_TEMPLATE_NOT_RECOGNIZED_MANUAL_REVIEW_REQUIRED')
 elif kind=='xlsx':
  z=zipfile.ZipFile(io.BytesIO(b));items=z.infolist()
  if len(items)>150 or sum(x.file_size for x in items)>3000000:raise ValueError('XLSX_EXPANSION_LIMIT')
  if any('vbaProject' in x.filename or 'externalLink' in x.filename for x in items):raise ValueError('ACTIVE_WORKBOOK_QUARANTINED')
  w=load_workbook(io.BytesIO(b),data_only=False,read_only=True);out=[]
  for s in w:
   if s.max_row>300 or s.max_column>10:raise ValueError('XLSX_DIMENSION_LIMIT')
   for i,row in enumerate(s.iter_rows(values_only=True),1):
    if i==1:continue
    if len(row)<7 or not row[0]:continue
    if any(isinstance(v,str) and v.startswith('=') for v in row):raise ValueError('SPREADSHEET_FORMULA_REVIEW_REQUIRED')
    field,val,unit,period,currency,ref,date=row[:7]
    out.append({'field':field,'rawValue':str(val),'period':period,'currency':currency,'sourceRef':ref,'sourceDate':date,'sourceLocation':{'sheet':s.title,'cell':f'B{i}'}})
 else:raise ValueError('UNSUPPORTED')
 print(json.dumps({'ok':True,'rows':out}))
except Exception as e:print(json.dumps({'ok':False,'code':str(e)[:180]}));sys.exit(1)

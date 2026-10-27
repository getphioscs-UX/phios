import pathlib,json,hashlib,subprocess,datetime,shutil,re
import openpyxl
from pypdf import PdfReader
R=pathlib.Path('C:/phios');D=R/'tools/review/hd-r6/import-20261027/r65';P=R/'.tmp/hd-r65-private';P.mkdir(parents=True,exist_ok=True)
sha=lambda p:hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()
save=lambda n,v:(D/n).write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
git=lambda *a:subprocess.check_output(['git',*a],cwd=R,text=True).strip()
hist=list((D.parent/'r64').rglob('*'))+list((R/'output/pdf/hd-r64-owner-review').glob('*.pdf'))
save('BASELINE.json',{'at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'head':git('rev-parse','HEAD'),'branch':git('branch','--show-current'),'status':git('status','--short'),'index':git('diff','--cached','--name-only'),'worktrees':git('worktree','list'),'freeDiskBytes':shutil.disk_usage(R).free,'protected':[{'path':str(p.relative_to(R)).replace('\\','/'),'sha256':sha(p)} for p in hist if p.is_file()]})
desktop=pathlib.Path('C:/Users/Guest Account/Desktop');names=['人类图总结.xlsx','Human Design Summary.xlsx','Thea Alexa Teng人类图解读报告书（更新）_.pdf','Terasa——职场成功守则全面报告书.pdf','PHIOS-HD-R65-REPORT-QUALITY-AND-SOURCE-AUDIT.md'];inventory=[];raw=[]
for name in names:
 p=desktop/name
 if not p.exists():inventory.append({'name':name,'status':'MISSING'});continue
 item={'name':name,'path':str(p),'bytes':p.stat().st_size,'sha256':sha(p),'status':'READABLE_PRIVATE_REVIEW_ONLY'};inventory.append(item)
 if p.suffix=='.xlsx':
  w=openpyxl.load_workbook(p,read_only=True,data_only=False);sheets=[];lang='en' if name.startswith('Human') else 'zh-Hans'
  for s in w:
   cells=[{'coordinate':c.coordinate,'value':c.value,'type':c.data_type} for row in s for c in row if c.value is not None]
   namespace='CORE_HD' if s.title in ['中心','闸门','通道','类型','内在权威','角色pg2','环境','PHS & RAVE'] else 'SEPARATE_OR_UNRESOLVED'
   if any(x in s.title for x in ['睡梦','Dream']):namespace='EXCLUDED_DREAMRAVE'
   if any(x in s.title for x in ['BG5','基因','星座','疗愈']):namespace='SPECIALIST_SEPARATE'
   sheets.append({'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'nonemptyCells':len(cells),'namespace':namespace,'rights':'UNESTABLISHED','semantic':'NOT_WHOLE_SHEET_APPROVED'})
   raw.append({'language':lang,'sourceSHA256':item['sha256'],'sheet':s.title,'namespace':namespace,'cells':cells})
  item['worksheets']=sheets;item['worksheetCount']=len(sheets)
 elif p.suffix=='.pdf':
  doc=PdfReader(p);texts=[pg.extract_text() or '' for pg in doc.pages];item.update({'pages':len(doc.pages),'pageSizes':sorted(set((round(float(pg.mediabox.width)),round(float(pg.mediabox.height))) for pg in doc.pages)),'privateCoverage':{k:[i+1 for i,t in enumerate(texts) if re.search(v,t,re.I)] for k,v in {'centers':'中心|Centers','channels':'通道|Channel','profile':'人生角色|Profile','health':'健康|医疗|healing|health','career':'职场|职业|BG5|career','dream':'睡梦|DreamRave'}.items()},'reuse':'COVERAGE_COMPARISON_ONLY_NO_PROSE_OR_FIGURE_COPY'})
  (P/(p.stem+'.txt')).write_text('\n\n'.join(f'PAGE {i+1}\n{t}' for i,t in enumerate(texts)),encoding='utf-8')
(P/'workbook-exact-cells.json').write_text(json.dumps(raw,ensure_ascii=False,default=str),encoding='utf-8')
save('SOURCE-INVENTORY.json',{'sources':inventory,'requestedVersion5Names':{n:(desktop/n).exists() for n in ['人类图总结(5).xlsx','Human Design Summary(5).xlsx']},'versionRule':'Supplied unsuffixed actual files identified by hash; do not claim byte identity to absent (5) sources.','privateExtractionPath':str(P),'rawContentPublished':False})
print(json.dumps({'head':git('rev-parse','HEAD'),'privateRawSheets':len(raw),'sources':[{'name':x['name'],'bytes':x.get('bytes'),'sheets':x.get('worksheetCount'),'pages':x.get('pages'),'status':x['status']} for x in inventory]},ensure_ascii=False))
